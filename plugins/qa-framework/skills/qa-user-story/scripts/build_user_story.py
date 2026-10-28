"""Build a single-sheet user story workbook (Portonics "User Story & Estimation Template") from a story JSON.

Usage:
    .venv/bin/python -I scripts/build_user_story.py <story.json> <out.xlsx>

The JSON shape is documented in reference/story-json.md; reference/example-story.json is a full example.
Image paths in the JSON are relative to the JSON file's folder (or absolute).
"""
import json
import os
import sys
from datetime import datetime

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "template", "user-story-template.xlsx")

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="FF345D7E")
FEATURE_FILL = PatternFill("solid", fgColor="FFF3F6F9")
THIN = Side(style="thin", color="FFB7C4D1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WIDTHS = {"A": 11, "B": 32, "C": 32, "D": 24, "E": 18, "F": 34, "G": 62, "H": 52, "I": 46, "J": 18, "K": 18}
FIRST, HEADER = 12, 11
LINE_PT = 13.5
MAX_IMG_PX = 520  # ~390 pt, so an image never pushes a row past Excel's 409-pt limit
# Column number -> JSON key for each story row (B:C is the merged Screen/MockUp cell).
COLUMNS = {1: "sl", 4: "feature", 5: "as_a", 6: "expectation", 7: "cos", 8: "ac", 9: "note", 10: "qa_comments", 11: "qa_requirements"}
HEADER_CELLS = {"project": "C5", "client": "C6", "created_by": "C7", "date_created": "E5",
                "business_lead": "E6", "tech_lead": "E7", "notes": "G5", "epic": "B9", "epic_id": "D9"}


def lines_needed(text, width):
    per_line = max(1, int(width * 1.2))
    return sum(max(1, -(-len(part) // per_line)) for part in str(text).split("\n"))


def main(json_path, out_path):
    with open(json_path, encoding="utf-8") as f:
        story = json.load(f)
    base = os.path.dirname(os.path.abspath(json_path))

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    ws.title = story.get("sheet_name", "User Story")[:31]

    header = story.get("header", {})
    for key, cell in HEADER_CELLS.items():
        if key in header:
            value = header[key]
            if key == "date_created" and isinstance(value, str):
                value = datetime.fromisoformat(value)
            ws[cell].value = value
    for key, cell in HEADER_CELLS.items():
        if key not in header and isinstance(ws[cell].value, str) and ws[cell].value.startswith("<"):
            ws[cell].value = None
    notes = ws[HEADER_CELLS["notes"]]
    is_link = isinstance(notes.value, str) and notes.value.startswith("http")
    notes.font = Font(name=FONT, size=10, color="FF0000FF" if is_link else "FF000000", underline="single" if is_link else None)

    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    for row in ws.iter_rows(min_row=1, max_row=9):
        for c in row:
            if isinstance(c, MergedCell) or c.coordinate == HEADER_CELLS["notes"]:
                continue
            f = c.font
            c.font = Font(name=FONT, size=10, bold=f.b, color=f.color)
            c.alignment = Alignment(wrap_text=True, vertical="center")
    ws["B9"].font = Font(name=FONT, size=12, bold=True)
    for r in (5, 6, 7):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[9].height = 30

    for col in range(1, 12):
        c = ws.cell(HEADER, col)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
    ws.row_dimensions[HEADER].height = 32

    rows = story["rows"]
    merged = {str(m) for m in ws.merged_cells.ranges}
    for i, item in enumerate(rows):
        r = FIRST + i
        if f"B{r}:C{r}" not in merged:
            ws.merge_cells(f"B{r}:C{r}")
        for col, key in COLUMNS.items():
            ws.cell(r, col).value = item.get(key) or None
        for col in range(1, 12):
            c = ws.cell(r, col)
            c.border = BOX
            if isinstance(c, MergedCell):
                continue
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if col == 1 else "left")
            c.font = Font(name=FONT, size=10, bold=col in (1, 4))
            if col in (1, 4, 5):
                c.fill = FEATURE_FILL

        text_lines = max(lines_needed(ws.cell(r, col).value or "", WIDTHS[letter])
                         for col, letter in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H"), (9, "I")))
        height = text_lines * LINE_PT + 12
        if item.get("image"):
            path = item["image"] if os.path.isabs(item["image"]) else os.path.join(base, item["image"])
            img = XLImage(path)
            box_px = int((WIDTHS["B"] + WIDTHS["C"]) * 7 + 10) - 16
            scale = min(box_px / img.width, MAX_IMG_PX / img.height)
            img.width, img.height = box_px, int(img.height * scale)
            img.anchor = OneCellAnchor(
                _from=AnchorMarker(col=1, colOff=pixels_to_EMU(8), row=r - 1, rowOff=pixels_to_EMU(8)),
                ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)),
            )
            ws.add_image(img)
            height = max(height, img.height * 0.75 + 14)
        if height > 409:
            print(f"WARNING: row {item.get('sl')} text needs {height:.0f} pt; Excel caps rows at 409 pt. Shorten or split the row.")
        ws.row_dimensions[r].height = min(409, height)
    for r in range(FIRST + len(rows), max(ws.max_row, 32) + 1):
        ws.row_dimensions[r].height = 14.25
        for col in range(1, 12):
            c = ws.cell(r, col)
            if not isinstance(c, MergedCell):
                c.value = None

    ws.freeze_panes = "D12"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out_path)
    print(f"saved {out_path}: {len(rows)} rows, {sum(1 for x in rows if x.get('image'))} images")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
